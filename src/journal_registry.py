"""Journal Registry & Ranking Module.

Responsibilities
----------------
* Load the journal registry (``data/journals.csv``) into ``Journal`` objects.
* Filter journals by the user's inclusion criteria and selected areas.
* Build ``journals.csv`` from a single *combined* ranking spreadsheet supplied
  by the user (CSV or XLSX) with flexible column-name matching.

The combined file is private (ABS/FT50/UTD24/ABDC/SJR lists are licensed) and is
therefore git-ignored once generated. See ``JOURNAL_SOURCES.md``.
"""
from __future__ import annotations

import argparse
import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

from .config_loader import (
    AREAS_PATH,
    DATA_DIR,
    JournalInclusion,
    parse_abs_level,
)

DEFAULT_REGISTRY_PATH = DATA_DIR / "journals.csv"

CANONICAL_COLUMNS = [
    "journal_name",
    "publisher",
    "areas",
    "issn",
    "abs_level",
    "ft50",
    "utd24",
    "abdc",
    "sjr",
    "impact_factor",
    "journal_weight",
]

# Journal-weight formula (see compute_journal_weight). Weights blend a ranking
# prestige score, an FT50/UTD24 elective score, and an impact-factor score.
IF_CAP = 15.0            # impact factor mapped to 1.0 at/above this value
W_RANK = 0.60            # weight of ranking prestige (ABS / ABDC)
W_ELECTIVE = 0.15        # weight of FT50 + UTD24 membership
W_IMPACT = 0.25          # weight of impact factor
assert abs(W_RANK + W_ELECTIVE + W_IMPACT - 1.0) < 1e-9

_QUARTILE_RANK = {"Q1": 1, "Q2": 2, "Q3": 3, "Q4": 4}
_TRUTHY = {"1", "true", "yes", "y", "x", "t"}


@dataclass
class Journal:
    journal_name: str
    publisher: str = ""
    areas: list[str] = field(default_factory=list)
    issn: list[str] = field(default_factory=list)
    abs_level: Optional[str] = None
    ft50: bool = False
    utd24: bool = False
    abdc: Optional[str] = None
    sjr: Optional[str] = None
    impact_factor: float = 0.0
    journal_weight: float = 0.0

    @property
    def abs_numeric(self) -> Optional[float]:
        return parse_abs_level(self.abs_level)


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def _split_multi(value: str) -> list[str]:
    if not value:
        return []
    parts = re.split(r"[;|]", value)
    return [p.strip() for p in parts if p.strip()]


def _to_bool(value: object) -> bool:
    return str(value).strip().lower() in _TRUTHY


def load_registry(path: Path | str = DEFAULT_REGISTRY_PATH) -> list[Journal]:
    """Load journals from a canonical ``journals.csv``."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(
            f"Journal registry not found: {path}. Build it from your combined "
            f"ranking file: python -m src.journal_registry --build <file>. "
            f"A sample is provided at data/journals.sample.csv."
        )
    journals: list[Journal] = []
    with path.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            name = (row.get("journal_name") or "").strip()
            if not name:
                continue
            try:
                weight = float(row.get("journal_weight") or 0)
            except ValueError:
                weight = 0.0
            try:
                impact = float(row.get("impact_factor") or 0)
            except ValueError:
                impact = 0.0
            journals.append(
                Journal(
                    journal_name=name,
                    publisher=(row.get("publisher") or "").strip(),
                    areas=_split_multi(row.get("areas", "")),
                    issn=_split_multi(row.get("issn", "")),
                    abs_level=(row.get("abs_level") or "").strip() or None,
                    ft50=_to_bool(row.get("ft50")),
                    utd24=_to_bool(row.get("utd24")),
                    abdc=(row.get("abdc") or "").strip() or None,
                    sjr=(row.get("sjr") or "").strip() or None,
                    impact_factor=impact,
                    journal_weight=weight,
                )
            )
    return journals


# --------------------------------------------------------------------------- #
# Filtering
# --------------------------------------------------------------------------- #
def _matches_area(journal: Journal, areas: Iterable[str]) -> bool:
    wanted = {a.strip().lower() for a in areas}
    if not wanted:
        return True
    return any(a.strip().lower() in wanted for a in journal.areas)


def _passes_inclusion(journal: Journal, crit: JournalInclusion) -> bool:
    """A journal is included if it satisfies ANY enabled ranking rule (union)."""
    if crit.abs_min_level is not None and journal.abs_numeric is not None:
        if journal.abs_numeric >= crit.abs_min_level:
            return True
    if crit.abdc_grades and journal.abdc:
        if journal.abdc.upper() in {g.upper() for g in crit.abdc_grades}:
            return True
    if crit.sjr_quartile and journal.sjr:
        want = _QUARTILE_RANK.get(crit.sjr_quartile.upper())
        have = _QUARTILE_RANK.get(journal.sjr.upper())
        if want is not None and have is not None and have <= want:
            return True
    if crit.include_ft50 and journal.ft50:
        return True
    if crit.include_utd24 and journal.utd24:
        return True
    return False


def is_elite(journal: Journal) -> bool:
    """Elite = FT50 or UTD24 or ABS/AJG 4*/4 or ABDC A*."""
    if journal.ft50 or journal.utd24:
        return True
    if journal.abs_numeric is not None and journal.abs_numeric >= 4.0:
        return True
    if journal.abdc and journal.abdc.upper() == "A*":
        return True
    return False


def filter_journals(
    journals: list[Journal],
    areas: Iterable[str],
    inclusion: JournalInclusion,
    elite_only: bool = False,
) -> list[Journal]:
    """Return journals matching a selected area AND a quality gate.

    The quality gate is either the elite rule (``elite_only=True``) or the
    configured inclusion criteria (union of ranking rules).
    """
    areas = list(areas)
    gate = is_elite if elite_only else (lambda j: _passes_inclusion(j, inclusion))
    return [j for j in journals if _matches_area(j, areas) and gate(j)]


def filter_elite(journals: list[Journal]) -> list[Journal]:
    """Return only the elite journals (no area filtering)."""
    return [j for j in journals if is_elite(j)]


# --------------------------------------------------------------------------- #
# Building journals.csv from a combined ranking file
# --------------------------------------------------------------------------- #
_HEADER_SYNONYMS = {
    "journal_name": {"journalname", "journal", "title", "journaltitle", "sourcetitle", "source", "name"},
    "publisher": {"publisher", "publishername"},
    "areas": {"area", "areas", "field", "fields", "discipline", "disciplines", "subject", "subjects", "category", "categories"},
    "issn": {"issn", "issns", "printissn", "eissn", "issnprint", "issnonline"},
    "abs_level": {"abs", "ajg", "abslevel", "ajglevel", "absrating", "ajgrating", "abs2021", "ajg2021", "abs2024"},
    "ft50": {"ft50", "ft", "financialtimes50"},
    "utd24": {"utd24", "utd", "utdallas24"},
    "abdc": {"abdc", "abdcrating", "abdcgrade", "abdc2022"},
    "sjr": {"sjr", "sjrquartile", "sjrbestquartile", "quartile", "bestquartile"},
    "impact_factor": {"impactfactor", "if", "jif", "journalimpactfactor",
                      "2yrmeancitedness", "twoyearmeancitedness", "meancitedness",
                      "citescore"},
    "journal_weight": {"journalweight", "weight"},
}


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (h or "").lower())


def _build_header_map(headers: list[str]) -> dict[str, str]:
    """Map source headers -> canonical column names."""
    mapping: dict[str, str] = {}
    issn_cols: list[str] = []
    for h in headers:
        n = _norm_header(h)
        for canon, syns in _HEADER_SYNONYMS.items():
            if n in syns or n == canon:
                if canon == "issn":
                    issn_cols.append(h)
                else:
                    mapping[h] = canon
                break
    # Remember all issn-like columns so we can merge print + electronic ISSNs.
    for c in issn_cols:
        mapping[c] = "issn"
    return mapping


def _rank_score(j: Journal) -> float:
    """Ranking prestige in [0, 1] from ABS/AJG, then ABDC, then list membership."""
    abs_num = j.abs_numeric
    if abs_num is not None:
        return min(1.0, abs_num / 4.5)          # 4* -> 1.0, 4 -> 0.889, 3 -> 0.667
    if j.abdc:
        return {"A*": 1.0, "A": 0.85, "B": 0.6, "C": 0.4}.get(j.abdc.upper(), 0.5)
    if j.ft50 or j.utd24:
        return 0.8
    return 0.5


def compute_journal_weight(j: Journal, if_cap: float = IF_CAP) -> float:
    """Blend ranking prestige, FT50/UTD24 membership and impact factor into 0..25.

    SJR is intentionally excluded. Impact factor is normalised against ``if_cap``
    (values at/above the cap score 1.0). The three sub-scores are combined with
    W_RANK / W_ELECTIVE / W_IMPACT and scaled to the 0..25 range used downstream.
    """
    rank = _rank_score(j)
    elective = (int(bool(j.ft50)) + int(bool(j.utd24))) / 2.0
    impact = min(1.0, (j.impact_factor or 0.0) / if_cap) if if_cap > 0 else 0.0
    combined = W_RANK * rank + W_ELECTIVE * elective + W_IMPACT * impact
    return round(25.0 * combined, 2)


def _read_rows(path: Path) -> tuple[list[str], list[dict]]:
    """Read CSV or XLSX rows as (headers, list-of-dicts)."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(
                "Reading .xlsx requires openpyxl: pip install openpyxl"
            ) from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h) if h is not None else "" for h in rows[0]]
        records = [
            {headers[i]: ("" if i >= len(r) or r[i] is None else r[i]) for i in range(len(headers))}
            for r in rows[1:]
        ]
        return headers, records
    # default: CSV
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        return list(reader.fieldnames or []), list(reader)


def build_registry_from_combined(
    source: Path | str,
    out_path: Path | str = DEFAULT_REGISTRY_PATH,
) -> list[Journal]:
    """Parse a combined ranking spreadsheet into canonical ``journals.csv``."""
    source = Path(source)
    out_path = Path(out_path)
    headers, rows = _read_rows(source)
    if not headers:
        raise ValueError(f"No data found in {source}")
    header_map = _build_header_map(headers)
    if "journal_name" not in header_map.values():
        raise ValueError(
            "Could not find a journal-name column in the combined file. "
            f"Headers seen: {headers}"
        )

    journals: list[Journal] = []
    for row in rows:
        agg: dict[str, list[str]] = {}
        for src_col, canon in header_map.items():
            val = row.get(src_col)
            if val is None:
                continue
            agg.setdefault(canon, []).append(str(val).strip())

        name = " ".join(agg.get("journal_name", [])).strip()
        if not name:
            continue
        issns: list[str] = []
        for chunk in agg.get("issn", []):
            issns.extend(_split_multi(chunk.replace(",", ";")))
        # de-dupe issns preserving order
        issns = list(dict.fromkeys(i for i in issns if i and i.lower() != "none"))

        j = Journal(
            journal_name=name,
            publisher=" ".join(agg.get("publisher", [])).strip(),
            areas=_split_multi(";".join(agg.get("areas", []))),
            issn=issns,
            abs_level=(agg.get("abs_level", [None])[0] or None),
            ft50=_to_bool(agg.get("ft50", ["0"])[0]),
            utd24=_to_bool(agg.get("utd24", ["0"])[0]),
            abdc=(agg.get("abdc", [None])[0] or None),
            sjr=(agg.get("sjr", [None])[0] or None),
        )
        try:
            j.impact_factor = float(agg.get("impact_factor", ["0"])[0] or 0)
        except ValueError:
            j.impact_factor = 0.0
        explicit_weight = agg.get("journal_weight", [""])[0]
        try:
            j.journal_weight = float(explicit_weight) if explicit_weight else compute_journal_weight(j)
        except ValueError:
            j.journal_weight = compute_journal_weight(j)
        journals.append(j)

    write_registry(journals, out_path)
    return journals


def write_registry(journals: list[Journal], out_path: Path | str = DEFAULT_REGISTRY_PATH) -> None:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(CANONICAL_COLUMNS)
        for j in journals:
            writer.writerow(
                [
                    j.journal_name,
                    j.publisher,
                    ";".join(j.areas),
                    ";".join(j.issn),
                    j.abs_level or "",
                    1 if j.ft50 else 0,
                    1 if j.utd24 else 0,
                    j.abdc or "",
                    j.sjr or "",
                    j.impact_factor,
                    j.journal_weight,
                ]
            )


def sync_areas(journals: list[Journal], areas_path: Path = AREAS_PATH) -> list[str]:
    """Append any new area labels found in the registry to ``areas.yaml``."""
    import yaml

    existing = []
    if areas_path.exists():
        existing = (yaml.safe_load(areas_path.read_text(encoding="utf-8")) or {}).get("areas", [])
    known = {a.strip().lower() for a in existing}
    added = []
    for j in journals:
        for a in j.areas:
            if a and a.strip().lower() not in known:
                existing.append(a)
                known.add(a.strip().lower())
                added.append(a)
    if added:
        areas_path.write_text(
            yaml.safe_dump({"areas": sorted(existing)}, allow_unicode=True, sort_keys=False),
            encoding="utf-8",
        )
    return added


def _cli() -> None:
    parser = argparse.ArgumentParser(description="Journal registry utilities.")
    parser.add_argument("--build", metavar="FILE", help="Combined ranking file (.csv/.xlsx) to import.")
    parser.add_argument("--out", default=str(DEFAULT_REGISTRY_PATH), help="Output journals.csv path.")
    parser.add_argument("--sync-areas", action="store_true", help="Append new areas to areas.yaml.")
    args = parser.parse_args()

    if args.build:
        journals = build_registry_from_combined(args.build, args.out)
        print(f"Wrote {len(journals)} journals to {args.out}")
        if args.sync_areas:
            added = sync_areas(journals)
            print(f"Added {len(added)} new area(s) to areas.yaml: {added}")
    else:
        journals = load_registry(args.out)
        print(f"Registry {args.out} contains {len(journals)} journals.")


if __name__ == "__main__":
    _cli()
