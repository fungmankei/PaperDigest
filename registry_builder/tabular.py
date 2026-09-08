"""Read CSV or XLSX files into (headers, list-of-dict-rows)."""
from __future__ import annotations

import csv
import re
from pathlib import Path


def norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (h or "").lower())


def read_table(path: Path) -> tuple[list[str], list[dict]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Reading .xlsx requires openpyxl: pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h) if h is not None else "" for h in rows[0]]
        records = [
            {headers[i]: ("" if i >= len(r) or r[i] is None else r[i]) for i in range(len(headers))}
            for r in rows[1:]
        ]
        return headers, records
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        return list(reader.fieldnames or []), list(reader)


def find_column(headers: list[str], candidates: set[str]) -> str | None:
    """Return the first header whose normalised form is in ``candidates``."""
    for h in headers:
        if norm_header(h) in candidates:
            return h
    return None


def read_matrix(path: Path) -> list[list]:
    """Read a CSV/XLSX as a raw list-of-rows (no header assumption)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Reading .xlsx requires openpyxl: pip install openpyxl") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [["" if c is None else c for c in row]
                for row in ws.iter_rows(values_only=True)]
    with path.open(encoding="utf-8-sig", newline="") as fh:
        return [list(row) for row in csv.reader(fh)]


def find_header_row(matrix: list[list], required: set[str]) -> int:
    """Index of the first row whose normalised cells include all ``required``."""
    for i, row in enumerate(matrix):
        norms = {norm_header(str(c)) for c in row}
        if required <= norms:
            return i
    return 0
